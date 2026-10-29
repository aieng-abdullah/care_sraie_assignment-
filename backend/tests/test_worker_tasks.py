"""Tests for Celery worker tasks: fan-out, OCR fallback, failure, finalization."""

from collections.abc import Iterator
from pathlib import Path
from uuid import UUID, uuid4

import pytest
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.core.database import get_session_factory
from app.models import Document, Page, ProcessingEvent
from app.models.enums import DocumentStatus
from app.pipelines import OcrResult
from app.services.file_validation import PDF, PNG, XLSX
from app.workers import tasks as worker_tasks
from app.workers.celery_app import celery_app

_PDF_TEXT = (
    "Invoice INV-777 issued to ACME Corp for professional services. "
    "Subtotal 500.00, tax 50.00, total 550.00. Payment due in 30 days. "
    "Purchase order PO-991 attached. Thank you. "
)


class FakeOcrEngine:
    """OCR double returning canned output; wired via monkeypatched factory."""

    def __init__(self) -> None:
        """Record calls for assertions."""
        self.calls = 0

    def recognize(self, image: "object") -> OcrResult:
        """Return fixed OCR output."""
        self.calls += 1
        return OcrResult(text="OCR FROM SCAN", confidence=0.8)


@pytest.fixture
def worker_env(db_env: Path, monkeypatch: pytest.MonkeyPatch) -> Iterator[Path]:
    """Run Celery tasks eagerly in-process (no broker, no worker process)."""
    celery_app.conf.task_always_eager = True
    celery_app.conf.task_store_eager_result = False
    yield db_env
    celery_app.conf.task_always_eager = False


@pytest.fixture(autouse=True)
def fake_ocr(monkeypatch: pytest.MonkeyPatch) -> FakeOcrEngine:
    """Replace the tesseract factory with a fake (no binary in CI/dev env)."""
    engine = FakeOcrEngine()
    monkeypatch.setattr(worker_tasks, "get_ocr_engine", lambda name: engine)
    return engine


def _make_text_pdf(path: Path, pages: int = 2) -> Path:
    """Write a real PDF with text layer on each page."""
    import pymupdf

    document = pymupdf.open()
    for _ in range(pages):
        page = document.new_page()
        page.insert_text((72, 72), _PDF_TEXT, fontsize=10)
    document.save(path)
    document.close()
    return path


def _make_document(session: Session, path: Path, content_type: str) -> Document:
    """Insert a queued document row pointing at an existing file."""
    document = Document(
        filename=path.name,
        content_type=content_type,
        original_sha256=uuid4().hex,
        size_bytes=path.stat().st_size,
        storage_path=str(path),
        status=DocumentStatus.QUEUED,
    )
    session.add(document)
    session.commit()
    return document


def _events(session: Session, document_id: UUID) -> list[str]:
    """Return lifecycle event names for a document, oldest first."""
    rows = session.scalars(
        select(ProcessingEvent)
        .where(ProcessingEvent.document_id == document_id)
        .order_by(ProcessingEvent.created_at)
    )
    return [row.event_type for row in rows]


def test_pdf_fanout_reaches_ok(worker_env: Path) -> None:
    """Two-page text PDF: pages stored, status ends ok, events recorded."""
    path = _make_text_pdf(worker_env / "doc.pdf", pages=2)
    session = get_session_factory()()
    try:
        document = _make_document(session, path, PDF)
        result = worker_tasks.process_document(str(document.id))
        assert result == "fanout:2"

        session.expire_all()
        refreshed = session.get(Document, document.id)
        assert refreshed is not None
        assert refreshed.status == DocumentStatus.OK
        assert refreshed.page_count == 2

        pages = list(session.scalars(select(Page).where(Page.document_id == document.id)))
        assert len(pages) == 2
        assert all("Invoice INV-777" in (p.text_content or "") for p in pages)
        assert all(p.source == "text_layer" for p in pages)

        events = _events(session, document.id)
        assert "started" in events
        assert events.count("page_done") == 2
        assert "done" in events
    finally:
        session.close()


def test_blank_pdf_uses_ocr_fallback(worker_env: Path, fake_ocr: FakeOcrEngine) -> None:
    """Scan-like PDF goes through the (faked) OCR engine with confidence stored."""
    import pymupdf

    path = worker_env / "scan.pdf"
    document_file = pymupdf.open()
    document_file.new_page()
    document_file.save(path)
    document_file.close()

    session = get_session_factory()()
    try:
        document = _make_document(session, path, PDF)
        worker_tasks.process_document(str(document.id))

        session.expire_all()
        refreshed = session.get(Document, document.id)
        assert refreshed is not None
        assert refreshed.status == DocumentStatus.OK
        page = session.scalar(select(Page).where(Page.document_id == document.id))
        assert page is not None
        assert page.source == "ocr"
        assert page.text_content == "OCR FROM SCAN"
        assert page.confidence is not None
        assert fake_ocr.calls == 1
    finally:
        session.close()


def test_page_limit_exceeded_marks_failed(
    worker_env: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """PDFs beyond MAX_PDF_PAGES fail with a clear reason (NFR hard cap)."""
    path = _make_text_pdf(worker_env / "big.pdf", pages=3)
    monkeypatch.setenv("MAX_PDF_PAGES", "2")
    get_settings.cache_clear()

    session = get_session_factory()()
    try:
        document = _make_document(session, path, PDF)
        result = worker_tasks.process_document(str(document.id))
        assert result == "failed"

        session.expire_all()
        refreshed = session.get(Document, document.id)
        assert refreshed is not None
        assert refreshed.status == DocumentStatus.FAILED
        assert refreshed.failure_reason is not None
        assert "limit is 2" in refreshed.failure_reason
        assert "failed" in _events(session, document.id)
    finally:
        get_settings.cache_clear()
        session.close()


def test_corrupted_pdf_fails_gracefully(worker_env: Path) -> None:
    """Unreadable file marks the document failed; the worker task itself returns."""
    path = worker_env / "broken.pdf"
    path.write_bytes(b"this is not a pdf at all")

    session = get_session_factory()()
    try:
        document = _make_document(session, path, PDF)
        result = worker_tasks.process_document(str(document.id))
        assert result == "failed"

        session.expire_all()
        refreshed = session.get(Document, document.id)
        assert refreshed is not None
        assert refreshed.status == DocumentStatus.FAILED
        assert refreshed.failure_reason
    finally:
        session.close()


def test_image_single_task_reaches_ok(worker_env: Path, fake_ocr: FakeOcrEngine) -> None:
    """PNG goes through process_single then finalizes to ok."""
    from PIL import Image

    path = worker_env / "photo.png"
    Image.new("RGB", (300, 200), "white").save(path)

    session = get_session_factory()()
    try:
        document = _make_document(session, path, PNG)
        result = worker_tasks.process_document(str(document.id))
        assert result == "single"

        session.expire_all()
        refreshed = session.get(Document, document.id)
        assert refreshed is not None
        assert refreshed.status == DocumentStatus.OK
        page = session.scalar(select(Page).where(Page.document_id == document.id))
        assert page is not None
        assert page.source == "ocr"
        assert fake_ocr.calls == 1
    finally:
        session.close()


def test_xlsx_single_task_stores_every_sheet(worker_env: Path) -> None:
    """Workbook sheets land as ordered page rows and status becomes ok."""
    from openpyxl import Workbook

    path = worker_env / "sheets.xlsx"
    book = Workbook()
    book.active.title = "A"
    book.active.append(["col1"])
    book.active.append(["v1"])
    book.create_sheet("B").append(["col2"])
    book.save(path)

    session = get_session_factory()()
    try:
        document = _make_document(session, path, XLSX)
        worker_tasks.process_document(str(document.id))

        session.expire_all()
        refreshed = session.get(Document, document.id)
        assert refreshed is not None
        assert refreshed.status == DocumentStatus.OK
        pages = list(
            session.scalars(
                select(Page).where(Page.document_id == document.id).order_by(Page.page_no)
            )
        )
        assert [p.page_no for p in pages] == [1, 2]
        assert "v1" in (pages[0].text_content or "")
    finally:
        session.close()


def test_finalize_is_idempotent(worker_env: Path) -> None:
    """Second finalize call on a terminal document is a no-op."""
    path = _make_text_pdf(worker_env / "once.pdf", pages=1)
    session = get_session_factory()()
    try:
        document = _make_document(session, path, PDF)
        worker_tasks.process_document(str(document.id))
        session.expire_all()
        assert session.get(Document, document.id) is not None
        first_status = session.get(Document, document.id)
        assert first_status is not None
        assert first_status.status == DocumentStatus.OK

        assert worker_tasks.finalize(str(document.id)) == "already-terminal"
        assert worker_tasks.finalize(str(document.id)) == "already-terminal"
    finally:
        session.close()


def test_dispatch_records_event_when_broker_down(
    worker_env: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Broker failure leaves the document queued and records enqueue_failed."""
    from app.services import documents as documents_service

    def _boom(document_id: str) -> None:
        raise ConnectionError("broker down")

    monkeypatch.setattr(worker_tasks.process_document, "delay", _boom)

    session = get_session_factory()()
    try:
        path = _make_text_pdf(worker_env / "q.pdf", pages=1)
        document = _make_document(session, path, PDF)
        ok = documents_service.dispatch_document(session, document.id)
        assert ok is False

        session.expire_all()
        refreshed = session.get(Document, document.id)
        assert refreshed is not None
        assert refreshed.status == DocumentStatus.QUEUED
        assert "enqueue_failed" in _events(session, document.id)
    finally:
        session.close()
